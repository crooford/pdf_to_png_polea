import os
from pathlib import Path
from docx import Document
from docx.shared import Inches
from config import settings
from utils.logger import logger

def create_word_docs():
    """
    Crea documentos de Word con las imágenes recortadas.
    """
    base_dir = settings.BASE_DIR
    images_per_row = settings.DOC_IMAGES_PER_ROW
    image_width = Inches(settings.DOC_IMAGE_WIDTH_INCHES)

    logger.info("Iniciando creación de documentos Word...")

    if not base_dir.exists():
        logger.error(f"Directorio base no encontrado: {base_dir}")
        return

    doc = Document()
    # Nota: El script original creaba un solo documento iterando sobre todo, 
    # pero guardaba el documento con un nombre diferente en cada iteración de imagen?
    # Revisando lógica original:
    # for category... for file (equipo)... for image...
    #   add to table
    #   output_docx = output_doc + "-" + category + "-" + file + ".docx"
    #   doc.save(output_docx)
    # Esto guardaba el documento MÚLTIPLES VECES, creciendo en tamaño cada vez.
    # Probablemente querían un documento POR EQUIPO o POR CATEGORÍA.
    # Aquí asumiremos UN DOCUMENTO POR EQUIPO para ser más limpio, o uno global si eso querían.
    # La lógica original acumulaba todo en 'doc' y luego guardaba con nombre diferente.
    # Vamos a cambiarlo a: Un documento por Equipo.

    for category_path in base_dir.iterdir():
        if not category_path.is_dir():
            continue

        category_name = category_path.name
        img_base_dir = category_path / "IMG"
        
        if not img_base_dir.exists():
            continue

        for team_dir in img_base_dir.iterdir():
            if not team_dir.is_dir():
                continue

            team_name = team_dir.name
            logger.info(f"Procesando equipo: {team_name}")
            
            # Recolectar imágenes recortadas
            cropped_images = sorted(
                [f for f in team_dir.iterdir() if f.name.startswith("r_") and f.suffix.lower() == ".png"],
                key=lambda x: x.name
            )

            if not cropped_images:
                logger.warning(f"No hay imágenes recortadas para {team_name}")
                continue

            # Crear un documento NUEVO para cada equipo
            team_doc = Document()
            team_doc.add_heading(f"Equipo: {team_name} ({category_name})", 0)

            # Calcular filas
            num_rows = (len(cropped_images) + images_per_row - 1) // images_per_row
            
            table = team_doc.add_table(rows=num_rows, cols=images_per_row)
            table.autofit = True

            for i, img_path in enumerate(cropped_images):
                row_idx = i // images_per_row
                col_idx = i % images_per_row
                
                cell = table.rows[row_idx].cells[col_idx]
                paragraph = cell.paragraphs[0]
                run = paragraph.add_run()
                try:
                    run.add_picture(str(img_path), width=image_width)
                except Exception as e:
                    logger.error(f"Error agregando imagen {img_path.name}: {e}")

            # Guardar documento
            output_filename = f"Ficha_{category_name}_{team_name}.docx"
            output_path = team_dir / output_filename # Guardar en la misma carpeta del equipo? O en base?
            # El original guardaba en output_doc + ...
            # Guardemos en la carpeta del equipo para mantener orden
            
            try:
                team_doc.save(str(output_path))
                logger.info(f"Documento guardado: {output_path}")
            except Exception as e:
                logger.error(f"Error guardando documento {output_filename}: {e}")

    logger.info("Creación de documentos finalizada.")
